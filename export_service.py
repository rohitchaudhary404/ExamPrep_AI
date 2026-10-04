#!/usr/bin/env python3
"""Export a study pack to Word (.docx) and PDF.

Uses python-docx for Word and fpdf2 for PDF. Unicode text is fully supported:
the PDF renderer registers a Unicode TrueType font (Segoe UI, then Arial, then
DejaVu) and only falls back to latin-1 sanitisation if no TTF font is found.
"""

import io
import os

_FONT_CANDIDATES = [
    {
        "family": "SegoeUI",
        "": r"C:\Windows\Fonts\segoeui.ttf",
        "B": r"C:\Windows\Fonts\segoeuib.ttf",
        "I": r"C:\Windows\Fonts\segoeuii.ttf",
        "BI": r"C:\Windows\Fonts\segoeuiz.ttf",
    },
    {
        "family": "Arial",
        "": r"C:\Windows\Fonts\arial.ttf",
        "B": r"C:\Windows\Fonts\arialbd.ttf",
        "I": r"C:\Windows\Fonts\ariali.ttf",
        "BI": r"C:\Windows\Fonts\arialbi.ttf",
    },
    {
        "family": "DejaVu",
        "": r"C:\Windows\Fonts\DejaVuSans.ttf",
        "B": r"C:\Windows\Fonts\DejaVuSans-Bold.ttf",
        "I": r"C:\Windows\Fonts\DejaVuSans-Oblique.ttf",
        "BI": r"C:\Windows\Fonts\DejaVuSans-BoldOblique.ttf",
    },
]

# used only when no Unicode TTF font can be registered (ASCII fallback)
_ASCII_MAP = {
    "—": "-", "–": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...",
    "→": "->", "←": "<-", "⇒": "=>", "•": "-", "·": "-", "²": "^2", "³": "^3",
    "°": " deg", "Ω": " ohm", "√": " sqrt", "×": "x", "÷": "/", "−": "-",
    "π": " pi", "θ": " theta", "α": " alpha", "β": " beta", "Δ": " delta",
    "₹": "Rs ", "±": "+/-", "≤": "<=", "≥": ">=", "≠": "!=",
}


def _sanitize_ascii(text):
    for k, v in _ASCII_MAP.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def _s(data, key, default=""):
    return str(data.get(key) or default)


def _context_meta(data):
    """Build the header meta line (handles school vs university context)."""
    ctx = data.get("_context") or {}
    level = (ctx.get("level") or "").lower()
    if level == "university":
        parts = [ctx.get("university") or "", ctx.get("course") or ""]
        sem = str(ctx.get("semester") or "")
        if sem:
            parts.append(f"Semester {sem}")
        parts.append(ctx.get("subject") or "")
        return " · ".join(p for p in parts if p) or "University"
    return " · ".join([
        f"Board: {ctx.get('board') or data.get('board') or '-'}",
        f"Class: {ctx.get('class_level') or data.get('class_level') or '-'}",
        f"Subject: {ctx.get('subject') or data.get('subject') or '-'}",
    ])


# --------------------------------------------------------------------------
# Word (.docx)
# --------------------------------------------------------------------------
def build_docx(data):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8F)
        return h

    def para(text, bold=False, size=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        return p

    # Title
    title = doc.add_heading("Study Pack — ExamPrep AI", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8F)

    meta = _context_meta(data)
    p = doc.add_paragraph()
    r = p.add_run(meta)
    r.italic = True
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    summary = _s(data, "summary").strip()
    if summary:
        heading("Summary", 1)
        para(summary)

    heading("Key Points", 1)
    for i, kp in enumerate(data.get("key_points") or [], 1):
        para(f"{i}. {kp}")

    heading("Important Topics", 1)
    for t in data.get("important_topics") or []:
        line = f"{_s(t, 'topic')}  —  {_s(t, 'weightage')}  ·  {_s(t, 'exam_probability')}"
        para(line, bold=True)
        if _s(t, "reason"):
            p = doc.add_paragraph()
            r = p.add_run(f"  {_s(t, 'reason')}")
            r.italic = True

    heading("Most Important Questions", 1)
    for i, q in enumerate(data.get("important_questions") or [], 1):
        marks = q.get("marks")
        mtag = f"  [{marks} marks]" if marks else ""
        qtype = _s(q, "type")
        ttag = f"  ({qtype})" if qtype else ""
        para(f"Q{i}. {_s(q, 'question')}{mtag}{ttag}")

    heading("Previous Year Questions (PYQs)", 1)
    pyqs = data.get("pyqs") or []
    if pyqs:
        for q in pyqs:
            head = f"[{q.get('year')}] {_s(q, 'subject')} (Class {q.get('class')})"
            para(head, bold=True)
            para(_s(q, "question"))
            if _s(q, "answer_hint"):
                p = doc.add_paragraph()
                r = p.add_run(f"Answer hint: {_s(q, 'answer_hint')}")
                r.italic = True
    else:
        para("No matching PYQs found in the local database.")

    web = data.get("web") or {}
    if web.get("enabled"):
        heading("Most Repeated PYQs (from the web)", 1)
        if _s(web, "trend"):
            para(_s(web, "trend"))
        for t in web.get("repeated_topics") or []:
            para(f"{_s(t, 'topic')}  —  {_s(t, 'frequency')}", bold=True)
        for q in web.get("repeated_questions") or []:
            para(_s(q, "question"))
            p = doc.add_paragraph()
            r = p.add_run(f"  {_s(q, 'frequency')}" + (f"  ·  {_s(q, 'type')}" if _s(q, "type") else ""))
            r.italic = True
            for u in (q.get("source_urls") or [])[:3]:
                ps = doc.add_paragraph()
                rs = ps.add_run(f"  {u}")
                rs.font.size = Pt(9)
                rs.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
class _StudyPDF:  # thin wrapper around FPDF to keep font config tidy
    def __init__(self):
        from fpdf import FPDF

        self.pdf = FPDF()
        self.family = None
        for cand in _FONT_CANDIDATES:
            family = cand["family"]
            regular = cand.get("")
            if regular and os.path.isfile(regular):
                for style in ("", "B", "I", "BI"):
                    fname = cand.get(style)
                    if fname and os.path.isfile(fname):
                        self.pdf.add_font(family, style, fname)
                self.family = family
                break

    def text(self, s):
        if self.family:
            return str(s)
        return _sanitize_ascii(str(s))


def build_pdf(data):
    doc = _StudyPDF()
    pdf = doc.pdf
    fam = doc.family or "helvetica"

    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_margins(14, 14, 14)
    LM = pdf.l_margin
    W = pdf.w - pdf.l_margin - pdf.r_margin  # usable text width

    def reset():
        pdf.set_x(LM)

    def section(title):
        pdf.ln(6)
        reset()
        pdf.set_font(fam, "B", 13)
        pdf.set_text_color(31, 58, 143)
        pdf.multi_cell(W, 7, doc.text(title))
        pdf.ln(1)
        reset()
        pdf.set_draw_color(200, 210, 230)
        pdf.line(LM, pdf.get_y(), LM + W, pdf.get_y())
        pdf.ln(3)
        reset()

    def body(text, size=10, style=""):
        reset()
        pdf.set_font(fam, style, size)
        pdf.set_text_color(30, 30, 30)
        pdf.multi_cell(W, 5.4, doc.text(text))

    # Title
    reset()
    pdf.set_font(fam, "B", 20)
    pdf.set_text_color(31, 58, 143)
    pdf.cell(W, 10, doc.text("Study Pack — ExamPrep AI"), align="C")
    pdf.ln(12)
    reset()
    pdf.set_font(fam, "", 10)
    pdf.set_text_color(110, 110, 110)
    meta = _context_meta(data)
    pdf.cell(W, 7, doc.text(meta), align="C")
    pdf.ln(4)
    reset()

    summary = _s(data, "summary").strip()
    if summary:
        section("Summary")
        body(summary)

    section("Key Points")
    for i, kp in enumerate(data.get("key_points") or [], 1):
        body(f"{i}. {kp}")

    section("Important Topics")
    for t in data.get("important_topics") or []:
        reset()
        pdf.set_font(fam, "B", 10.5)
        pdf.set_text_color(20, 20, 20)
        pdf.multi_cell(W, 5.4, doc.text(
            f"{_s(t, 'topic')}  —  {_s(t, 'weightage')}  ·  {_s(t, 'exam_probability')}"))
        if _s(t, "reason"):
            reset()
            pdf.set_font(fam, "I", 9.5)
            pdf.set_text_color(90, 90, 90)
            pdf.multi_cell(W, 5, doc.text("  " + _s(t, "reason")))
        pdf.ln(1.5)
        reset()

    section("Most Important Questions")
    for i, q in enumerate(data.get("important_questions") or [], 1):
        marks = q.get("marks")
        mtag = f"  [{marks} marks]" if marks else ""
        qtype = _s(q, "type")
        ttag = f"  ({qtype})" if qtype else ""
        body(f"Q{i}. {_s(q, 'question')}{mtag}{ttag}")

    section("Previous Year Questions (PYQs)")
    pyqs = data.get("pyqs") or []
    if pyqs:
        for q in pyqs:
            reset()
            pdf.set_font(fam, "B", 9.5)
            pdf.set_text_color(20, 20, 20)
            pdf.multi_cell(W, 5, doc.text(
                f"[{q.get('year')}] {_s(q, 'subject')} (Class {q.get('class')})"))
            body(_s(q, "question"), size=10)
            if _s(q, "answer_hint"):
                reset()
                pdf.set_font(fam, "I", 9)
                pdf.set_text_color(90, 90, 90)
                pdf.multi_cell(W, 5, doc.text("Answer hint: " + _s(q, "answer_hint")))
            pdf.ln(2)
            reset()
    else:
        body("No matching PYQs found in the local database.")

    web = data.get("web") or {}
    if web.get("enabled"):
        section("Most Repeated PYQs (from the web)")
        if _s(web, "trend"):
            body(_s(web, "trend"))
            pdf.ln(2)
        for t in web.get("repeated_topics") or []:
            reset()
            pdf.set_font(fam, "B", 10)
            pdf.set_text_color(20, 20, 20)
            pdf.multi_cell(W, 5.4, doc.text(f"{_s(t, 'topic')}  —  {_s(t, 'frequency')}"))
        pdf.ln(1)
        for q in web.get("repeated_questions") or []:
            body(_s(q, "question"))
            reset()
            pdf.set_font(fam, "I", 9)
            pdf.set_text_color(90, 90, 90)
            pdf.multi_cell(W, 5, doc.text("  " + _s(q, "frequency") + (f"  ·  {_s(q, 'type')}" if _s(q, "type") else "")))
            for u in (q.get("source_urls") or [])[:3]:
                reset()
                pdf.set_font(fam, "", 8)
                pdf.set_text_color(59, 130, 246)
                pdf.multi_cell(W, 4, doc.text("  " + u))
            pdf.ln(1.5)

    # footer note
    pdf.ln(4)
    reset()
    pdf.set_font(fam, "I", 8)
    pdf.set_text_color(150, 150, 150)
    pdf.multi_cell(W, 4, doc.text("Generated by ExamPrep AI."))

    return bytes(pdf.output())
